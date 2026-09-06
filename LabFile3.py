from docx import Document
import pandas as pd
import matplotlib.pyplot as plt

# ============================================================
# LOAD WORD FILE
# ============================================================

file_path = r"C:\Users\Rudra\OneDrive\Documents\Design and analysis files\lab_file.docx"

doc = Document(file_path)

print("Number of tables:", len(doc.tables))


# ============================================================
# FUNCTION TO CONVERT WORD TABLE TO PANDAS DATAFRAME
# ============================================================

def read_table(table):
    data = []

    for row in table.rows:
        values = [cell.text.strip() for cell in row.cells]

        # Ignore completely empty rows
        if any(values):
            data.append(values)

    df = pd.DataFrame(data[1:], columns=data[0])

    return df
  # ============================================================
# READ ALL 6 TABLES
# ============================================================

tables = []

for i, table in enumerate(doc.tables):
    df = read_table(table)
    tables.append(df)

    print(f"\nTable {i + 1}:")
    print(df)


# TABLES 1, 2, 3 = EXECUTION TIME
# TABLES 4, 5, 6 = MEMORY CONSUMPTION
# ============================================================

time_tables = tables[0:3]
memory_tables = tables[3:6]

# ============================================================
# CONVERT EXECUTION TIME FROM STRING TO FLOAT
# ============================================================

for df in time_tables:

    for column in df.columns[1:]:

        df[column] = (
            df[column]
            .astype(str)
            .str.replace(" sec", "", regex=False)
            .astype(float)
        )

    df["Input Size"] = pd.to_numeric(df["Input Size"])

# ============================================================
# CONVERT MEMORY FROM STRING TO FLOAT
# ============================================================

for df in memory_tables:

    for column in df.columns[1:]:

        df[column] = (
            df[column]
            .astype(str)
            .str.replace("KB", "", regex=False)
            .str.strip()
            .astype(float)
        )

    df["Input Size"] = pd.to_numeric(df["Input Size"])

# ============================================================
# CALCULATE AVERAGE EXECUTION TIME
# ============================================================

time_df = time_tables[0].copy()

algorithms = [
    "Bubble Sort",
    "Selection Sort",
    "Insertion Sort",
    "Merge Sort",
    "Quick Sort"
]

for algorithm in algorithms:

    time_df[algorithm] = (
        time_tables[0][algorithm]
        + time_tables[1][algorithm]
        + time_tables[2][algorithm]
    ) / 3


# ============================================================
# CONVERT MEMORY TABLES TO NUMERIC
# ============================================================

for df in memory_tables:

    # Convert Input Size
    df["Input Size"] = pd.to_numeric(
        df["Input Size"],
        errors="coerce"
    )

    # Convert memory columns
    for column in df.columns[1:]:

        df[column] = (
            df[column]
            .astype(str)
            .str.replace("KB", "", regex=False)
            .str.strip()
        )

        df[column] = pd.to_numeric(
            df[column],
            errors="coerce"
        )


# ============================================================
# CALCULATE AVERAGE MEMORY CONSUMPTION
# ============================================================

memory_df = memory_tables[0].copy()

for algorithm in algorithms:

    memory_df[algorithm] = (
        memory_tables[0][algorithm]
        + memory_tables[1][algorithm]
        + memory_tables[2][algorithm]
    ) / 3


# ============================================================
# CHECK DATA TYPES
# ============================================================

print("\nExecution Time Data Types:")
print(time_df.dtypes)

print("\nMemory Data Types:")
print(memory_df.dtypes)

# ============================================================
# SAVE PROCESSED DATA TO CSV
# ============================================================

time_df.to_csv("average_execution_time.csv", index=False)

memory_df.to_csv("average_memory_consumption.csv", index=False)

# ============================================================
# GRAPH 1: EXECUTION TIME
# ============================================================

plt.figure(figsize=(10, 6))

for algorithm in algorithms:

    plt.plot(
        time_df["Input Size"],
        time_df[algorithm],
        marker="o",
        linewidth=2,
        label=algorithm
    )

plt.title("Sorting Algorithms - Average Execution Time")
plt.xlabel("Input Size")
plt.ylabel("Execution Time (seconds)")
plt.legend()
plt.grid(True, linestyle="--", alpha=0.5)

plt.tight_layout()

plt.savefig(
    "sorting_execution_time.png",
    dpi=300
)

plt.show()

# ============================================================
# GRAPH 2: MEMORY CONSUMPTION
# ============================================================

plt.figure(figsize=(10, 6))

for algorithm in algorithms:

    plt.plot(
        memory_df["Input Size"],
        memory_df[algorithm],
        marker="o",
        linewidth=2,
        label=algorithm
    )

plt.title("Sorting Algorithms - Average Memory Consumption")
plt.xlabel("Input Size")
plt.ylabel("Memory Consumption (KB)")
plt.legend()
plt.grid(True, linestyle="--", alpha=0.5)

plt.tight_layout()

plt.savefig(
    "sorting_memory_consumption.png",
    dpi=300
)

plt.show()


# ============================================================
# GRAPH 3: EXECUTION TIME - BAR CHART FOR 5000 INPUT
# ============================================================

last_row = time_df.iloc[-1]

plt.figure(figsize=(10, 6))

plt.bar(
    algorithms,
    [last_row[algorithm] for algorithm in algorithms],
    color=[
        "red",
        "orange",
        "green",
        "blue",
        "purple"
    ]
)

plt.title("Execution Time Comparison for Input Size 5000")
plt.xlabel("Sorting Algorithm")
plt.ylabel("Execution Time (seconds)")
plt.xticks(rotation=20)
plt.grid(axis="y", linestyle="--", alpha=0.5)

plt.tight_layout()

plt.savefig(
    "execution_time_5000.png",
    dpi=300
)

plt.show()

# ============================================================
# GRAPH 4: MEMORY COMPARISON FOR 5000 INPUT
# ============================================================

last_memory = memory_df.iloc[-1]

plt.figure(figsize=(10, 6))

plt.bar(
    algorithms,
    [last_memory[algorithm] for algorithm in algorithms],
    color=[
        "red",
        "orange",
        "green",
        "blue",
        "purple"
    ]
)

plt.title("Memory Consumption Comparison for Input Size 5000")
plt.xlabel("Sorting Algorithm")
plt.ylabel("Memory Consumption (KB)")
plt.xticks(rotation=20)
plt.grid(axis="y", linestyle="--", alpha=0.5)

plt.tight_layout()

plt.savefig(
    "memory_5000.png",
    dpi=300
)

plt.show()

# ============================================================
# OBSERVATIONS
# ============================================================

print("\n")
print("=" * 70)
print("OBSERVATIONS")
print("=" * 70)

print("""
1. Execution Time:
   Bubble Sort, Selection Sort and Insertion Sort show a
   significant increase in execution time as the input size
   increases.

2. Selection Sort:
   Selection Sort has the highest execution time for large
   inputs. For an input size of 5000, its average execution
   time is much higher than Merge Sort and Quick Sort.

3. Bubble Sort:
   Bubble Sort also becomes considerably slower for large
   input sizes because of its O(n²) time complexity.

4. Insertion Sort:
   Insertion Sort performs well for small inputs but its
   execution time increases rapidly for larger inputs.

5. Merge Sort:
   Merge Sort performs significantly better than the
   quadratic sorting algorithms for large input sizes.
   Its average time complexity is O(n log n).

6. Quick Sort:
   Quick Sort generally gives the best execution-time
   performance for the larger input sizes in these results.

7. Memory Consumption:
   Merge Sort and Quick Sort generally consume more memory
   than Bubble Sort and Insertion Sort because of additional
   data structures and recursive operations.

8. Overall:
   The results show that algorithms with better asymptotic
   time complexity are more suitable for processing large
   datasets.

9. The three runs were averaged before plotting, which gives
   a more reliable comparison and reduces the effect of
   variations between individual runs.
""")

# ============================================================
# PRINT FINAL COMPARISON AT INPUT SIZE 5000
# ============================================================

print("\n")
print("=" * 70)
print("FINAL COMPARISON FOR INPUT SIZE 5000")
print("=" * 70)

for algorithm in algorithms:

    print(
        f"{algorithm:<20} "
        f"Time: {time_df.iloc[-1][algorithm]:.6f} sec   "
        f"Memory: {memory_df.iloc[-1][algorithm]:.2f} KB"
    )
